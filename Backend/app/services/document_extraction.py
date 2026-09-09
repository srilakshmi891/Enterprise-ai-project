# app/services/document_extraction.py
"""Document extraction service.
Extracts plain text from supported file types: PDF, DOCX, TXT, MD.
Raises ExtractionError on failure.
"""

import os
from typing import Any

from pypdf import PdfReader
from docx import Document as DocxDocument

class ExtractionError(Exception):
    """Raised when extraction fails for any reason."""
    pass


def _extract_pdf(file_path: str) -> str:
    try:
        reader = PdfReader(file_path)
        text_parts = []
        for page in reader.pages:
            text = page.extract_text() or ""
            text_parts.append(text)
        combined = "\n".join(text_parts).strip()
        if not combined:
            raise ExtractionError("PDF contains no extractable text")
        return combined
    except Exception as e:
        raise ExtractionError(f"Failed to extract PDF: {e}")



def _extract_docx(file_path: str) -> str:
    try:
        doc = DocxDocument(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text]
        # Extract tables
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text for cell in row.cells if cell.text]
                if cells:
                    table_texts.append(" ".join(cells))
        all_parts = paragraphs + table_texts
        combined = "\n".join(all_parts).strip()
        if not combined:
            raise ExtractionError("DOCX contains no extractable text")
        return combined
    except Exception as e:
        raise ExtractionError(f"Failed to extract DOCX: {e}")


def _extract_txt(file_path: str) -> str:
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read().strip()
    except UnicodeDecodeError:
        # fallback to latin-1
        try:
            with open(file_path, "r", encoding="latin-1") as f:
                return f.read().strip()
        except Exception as e:
            raise ExtractionError(f"Failed to read text file: {e}")
    except Exception as e:
        raise ExtractionError(f"Failed to read text file: {e}")


def extract_text(file_path: str, file_type: str) -> str:
    """Extract text based on file_type.
    Supported file_type values: pdf, docx, txt, md.
    """
    if not os.path.isfile(file_path):
        raise ExtractionError("File does not exist on server")
    ft = file_type.lower()
    if ft == "pdf":
        return _extract_pdf(file_path)
    if ft == "docx":
        return _extract_docx(file_path)
    if ft in {"txt", "md"}:
        return _extract_txt(file_path)
    raise ExtractionError(f"Unsupported file_type for extraction: {file_type}")
